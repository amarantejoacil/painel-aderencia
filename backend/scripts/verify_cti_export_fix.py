# -*- coding: utf-8 -*-
"""Verificação pós-correção CTI/PAC — somente leitura + export simulado."""
from __future__ import annotations

import json
from io import BytesIO

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import IssueAuditQuery
from app.services.azure_devops_config import get_settings_row, resolve_pat, to_view
from app.services.issue_audit.client import AzureCollectionClient
from app.services.issue_audit.excel_export import build_audit_workbook
from app.services.issue_audit.runner import collect_comments


def excel_full_rows(result: dict, query: dict, issue_id: int, base_url: str) -> int:
    payload = dict(result)
    issue = next(i for i in payload["issues"] if int(i["id"]) == issue_id)
    payload["issues"] = [issue]
    buffer, _ = build_audit_workbook(
        query,
        payload,
        base_url=base_url,
        include_all_discussions=True,
    )
    wb = load_workbook(BytesIO(buffer.getvalue()))
    sheet = wb["Discussões completas"]
    return sum(1 for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] == issue_id)


def main() -> None:
    db = SessionLocal()
    view = to_view(get_settings_row(db))
    pat = resolve_pat(get_settings_row(db))
    row_cti = db.get(IssueAuditQuery, 3)
    row_pac = db.get(IssueAuditQuery, 11)
    cti = AzureCollectionClient(base_url=view.base_url, organization="CTI", pat=pat)
    pac = AzureCollectionClient(base_url=view.base_url, organization="Nucleo", pat=pat)
    out = []
    for label, client, row, strategy, samples in (
        ("CTI", cti, row_cti, "history", [1667, 1260]),
        ("PAC", pac, row_pac, "comments_api", [26689, 26673]),
    ):
        issues = {int(i["id"]): i for i in (row.result or {}).get("issues") or []}
        query = {
            "organization": row.organization,
            "project": row.project,
            "work_item_type": row.work_item_type,
            "period_criterion": row.period_criterion,
            "year": row.year,
            "comment_scope": row.comment_scope,
            "person_name": row.person_name,
            "person_identity_id": row.person_identity_id,
            "profile_id": row.profile_id,
        }
        for iid in samples:
            live = collect_comments(client, project=row.project, work_item_id=iid, strategy=strategy)
            stored = issues.get(iid) or {}
            all_comments = stored.get("all_comments")
            excel_rows = None
            if all_comments is not None:
                excel_rows = excel_full_rows(row.result, query, iid, view.base_url)
            ui = cti.fetch_ui_comment_counts([iid]).get(iid) if label == "CTI" else None
            api_probe = (
                cti.probe_project_comments_api(row.project, iid)
                if label == "CTI"
                else pac.probe_project_comments_api(row.project, iid)
            )
            out.append(
                {
                    "projeto": label,
                    "issue": iid,
                    "reextract": len(live),
                    "stored_all_comments": len(all_comments) if all_comments is not None else None,
                    "excel_full_rows": excel_rows,
                    "azure_ui_count": ui,
                    "project_api_count": api_probe.get("count") if api_probe.get("available") else None,
                }
            )
    print(json.dumps(out, ensure_ascii=False, indent=2))
    cti.close()
    pac.close()
    db.close()


if __name__ == "__main__":
    main()
