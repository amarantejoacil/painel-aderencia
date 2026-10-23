# -*- coding: utf-8 -*-
"""Validação por amostragem — somente leitura."""
from __future__ import annotations

import json
from collections import Counter
from io import BytesIO

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import IssueAuditQuery
from app.services.azure_devops_config import get_settings_row, resolve_pat, to_view
from app.services.issue_audit.client import AzureCollectionClient
from app.services.issue_audit.excel_export import build_audit_workbook
from app.services.issue_audit.matching import PersonTarget, person_matches, normalize_name
from app.services.issue_audit.runner import collect_comments

MARCIA = "Marcia Regina De Carvalho Buhr"
NUCLEO_ID = "9982c2e2-276e-480d-b049-a29c7d1a80e5"
CTI_ID = "1e66179c-7ef1-40b6-8533-97418539f742"


def load_query(db, qid: int):
    row = db.get(IssueAuditQuery, qid)
    if not row or not row.result:
        raise SystemExit(f"Consulta {qid} indisponível.")
    issues = {int(i["id"]): i for i in row.result.get("issues") or [] if i.get("id") is not None}
    return row, issues


def independent_fetch(client, project: str, issue_id: int, strategy: str) -> dict:
    item = client.request(
        "GET",
        f"{client.org_base}/_apis/wit/workitems/{issue_id}?fields=System.CommentCount,System.TeamProject,System.Title&api-version=7.1",
    )
    fields = item.get("fields") or {}
    try:
        comments = collect_comments(client, project=project, work_item_id=issue_id, strategy=strategy)
        fetch_err = None
    except Exception as exc:  # noqa: BLE001
        comments = []
        fetch_err = str(exc)
    alt_strategy = "history" if strategy == "comments_api" else "comments_api"
    try:
        alt = collect_comments(client, project=project, work_item_id=issue_id, strategy=alt_strategy)
        alt_err = None
    except Exception as exc:  # noqa: BLE001
        alt = []
        alt_err = str(exc)
    return {
        "ui_comment_count": fields.get("System.CommentCount"),
        "team_project": fields.get("System.TeamProject"),
        "title": fields.get("System.Title"),
        "independent_count": len(comments),
        "independent_fetch_error": fetch_err,
        "alt_strategy": alt_strategy,
        "alt_count": len(alt),
        "alt_error": alt_err,
        "comments": comments,
    }


def marcia_comments(comments: list[dict], identity_id: str) -> list[dict]:
    person = PersonTarget(MARCIA, identity_id=identity_id)
    return [c for c in comments if person_matches(c, person)]


def count_status(stored, ui, reextract) -> tuple[str, str]:
    if reextract is None:
        return "Inconclusivo", "Reextração independente falhou."
    if ui is not None and stored == reextract == ui:
        return "Conforme", "Sistema, System.CommentCount e reextração coincidem."
    if ui is not None and reextract == ui and stored != reextract:
        return "Divergente", f"Sistema={stored}, Azure UI={ui}, reextração={reextract}."
    if ui is not None and reextract != ui:
        return "Divergente", f"Sistema={stored}, Azure UI={ui}, reextração={reextract}."
    if stored == reextract:
        return "Conforme", f"Sistema={stored}, reextração={reextract}; System.CommentCount indisponível."
    return "Divergente", f"Sistema={stored}, reextração={reextract}."


def main() -> None:
    db = SessionLocal()
    view = to_view(get_settings_row(db))
    pat = resolve_pat(get_settings_row(db))
    row_cti, issues_cti = load_query(db, 3)
    row_pac, issues_pac = load_query(db, 11)
    cti_client = AzureCollectionClient(base_url=view.base_url, organization="CTI", pat=pat)
    pac_client = AzureCollectionClient(base_url=view.base_url, organization="Nucleo", pat=pat)

    cti_ok = [i for i in issues_cti.values() if not i.get("error")]
    pac_ok = [i for i in issues_pac.values() if not i.get("error")]

    cti_count_samples = sorted(
        [i for i in cti_ok if (i.get("comment_count") or 0) > 0],
        key=lambda x: x.get("comment_count") or 0,
        reverse=True,
    )[:2]
    cti_zero = next((i for i in cti_ok if (i.get("comment_count") or 0) == 0), None)
    if cti_zero:
        cti_count_samples.append(cti_zero)

    pac_count_ids = [26689]
    for issue in sorted(pac_ok, key=lambda x: x.get("comment_count") or 0, reverse=True):
        if int(issue["id"]) != 26689 and len(pac_count_ids) < 3:
            pac_count_ids.append(int(issue["id"]))

    cti_with = [i for i in cti_ok if i.get("person_commented")][:3]
    pac_with = [i for i in pac_ok if i.get("person_commented")][:3]
    if not any(int(i["id"]) == 26689 for i in pac_with):
        pac_with = [issues_pac[26689]] + [i for i in pac_with if int(i["id"]) != 26689][:2]

    cti_without = [i for i in cti_ok if not i.get("person_commented")][:3]
    pac_without = [i for i in pac_ok if not i.get("person_commented")][:3]

    cti_multi = sorted([i for i in cti_ok if (i.get("comment_count") or 0) >= 5], key=lambda x: x.get("comment_count") or 0, reverse=True)[:2]
    pac_multi = sorted([i for i in pac_ok if (i.get("comment_count") or 0) >= 10], key=lambda x: x.get("comment_count") or 0, reverse=True)[:2]
    if 26689 not in [int(i["id"]) for i in pac_multi]:
        pac_multi = [issues_pac[26689]] + pac_multi[:1]

    out = {
        "count_tests": [],
        "authorship_tests": [],
        "false_negative_tests": [],
        "integrity_tests": [],
        "excel_tests": [],
        "filter_tests": {},
    }

    for label, client, project, strategy, identity_id, issue in [
        *[("CTI", cti_client, row_cti.project, "history", CTI_ID, i) for i in cti_count_samples],
        *[("PAC", pac_client, row_pac.project, "comments_api", NUCLEO_ID, issues_pac[iid]) for iid in pac_count_ids],
    ]:
        iid = int(issue["id"])
        ind = independent_fetch(client, project, iid, strategy)
        status, div = count_status(issue.get("comment_count"), ind["ui_comment_count"], ind["independent_count"])
        out["count_tests"].append(
            {
                "teste": "Quantidade",
                "projeto": label,
                "issue": iid,
                "resultado": status,
                "divergencia": div,
                "azure_ui": ind["ui_comment_count"],
                "sistema": issue.get("comment_count"),
                "reextract": ind["independent_count"],
                "fonte_independente": strategy,
            }
        )

    for label, client, project, strategy, identity_id, samples in [
        ("CTI", cti_client, row_cti.project, "history", CTI_ID, cti_with),
        ("PAC", pac_client, row_pac.project, "comments_api", NUCLEO_ID, pac_with),
    ]:
        for issue in samples:
            iid = int(issue["id"])
            ind = independent_fetch(client, project, iid, strategy)
            person = marcia_comments(ind["comments"], identity_id)
            stored_person = int(issue.get("person_comment_count") or 0)
            if stored_person != len(person):
                status, div = "Divergente", f"Sistema={stored_person}, reextração={len(person)} comentários da Marcia."
            elif stored_person == 0:
                status, div = "Divergente", "Sistema marcou comentário da Marcia, mas reextração não encontrou autoria."
            else:
                missing = 0
                for sc in issue.get("comments") or []:
                    ok = any(
                        (c.get("id") and sc.get("id") and c.get("id") == sc.get("id"))
                        or (
                            normalize_name(c.get("author")) == normalize_name(sc.get("author"))
                            and (c.get("text") or "")[:120] == (sc.get("text") or "")[:120]
                        )
                        for c in person
                    )
                    if not ok:
                        missing += 1
                if missing:
                    status, div = "Divergente", f"{missing} comentário(s) armazenado(s) da Marcia sem par na reextração."
                else:
                    sample = person[0]
                    status = "Conforme"
                    div = (
                        f"{stored_person} comentário(s) confirmados; exemplo autor={sample.get('author')}, "
                        f"id={sample.get('author_id')}, data={sample.get('created_at_local') or sample.get('created_at')}."
                    )
            out["authorship_tests"].append(
                {"teste": "Autoria", "projeto": label, "issue": iid, "resultado": status, "divergencia": div}
            )

    for label, client, project, strategy, identity_id, samples in [
        ("CTI", cti_client, row_cti.project, "history", CTI_ID, cti_without),
        ("PAC", pac_client, row_pac.project, "comments_api", NUCLEO_ID, pac_without),
    ]:
        for issue in samples:
            iid = int(issue["id"])
            ind = independent_fetch(client, project, iid, strategy)
            person = marcia_comments(ind["comments"], identity_id)
            stored_person = int(issue.get("person_comment_count") or 0)
            if person and stored_person == 0:
                status, div = "Divergente", f"Possível falso negativo: reextração encontrou {len(person)} comentário(s) da Marcia."
            elif not person and stored_person == 0:
                status, div = "Conforme", "Reextração independente também não encontrou comentários da Marcia."
            else:
                status, div = "Divergente", f"Sistema marcou {stored_person}, reextração encontrou {len(person)}."
            out["false_negative_tests"].append(
                {"teste": "Falso negativo", "projeto": label, "issue": iid, "resultado": status, "divergencia": div}
            )

    for label, client, project, strategy, row, samples in [
        ("CTI", cti_client, row_cti.project, "history", row_cti, cti_multi),
        ("PAC", pac_client, row_pac.project, "comments_api", row_pac, pac_multi),
    ]:
        for issue in samples:
            iid = int(issue["id"])
            ind = independent_fetch(client, project, iid, strategy)
            comments = ind["comments"]
            ids = [c.get("id") for c in comments if c.get("id")]
            dup = [k for k, v in Counter(ids).items() if v > 1]
            empty = sum(1 for c in comments if not str(c.get("text") or "").strip())
            stored_all = issue.get("all_comments") if issue.get("all_comments") is not None else issue.get("comments")
            if issue.get("comment_count") != len(comments):
                status = "Divergente"
                div = f"Contagem divergente: sistema={issue.get('comment_count')}, reextração={len(comments)}."
            elif dup or empty:
                status = "Divergente"
                div = f"Duplicados={dup[:5]}; vazios={empty}."
            else:
                status = "Conforme"
                div = f"{len(comments)} comentários, {len({c.get('author_id') or c.get('author') for c in comments})} autores, sem duplicidade."
            out["integrity_tests"].append(
                {
                    "teste": "Integridade",
                    "projeto": label,
                    "issue": iid,
                    "resultado": status,
                    "divergencia": div,
                    "all_comments_len": len(stored_all or []),
                }
            )
            query = {
                "organization": row.organization,
                "project": row.project,
                "work_item_type": row.work_item_type,
                "period_criterion": row.period_criterion,
                "year": row.year,
                "comment_scope": row.comment_scope,
                "person_name": row.person_name,
                "person_identity_id": row.person_identity_id,
                "profile_id": row.profile_id,
            }
            payload = dict(row.result or {})
            payload["issues"] = [issue]
            buffer, _ = build_audit_workbook(query, payload, base_url=view.base_url, include_all_discussions=True)
            wb = load_workbook(BytesIO(buffer.getvalue()))
            sheet = wb["Discussões completas"]
            excel_rows = [r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == iid]
            if len(excel_rows) != len(comments):
                excel_status = "Divergente"
                excel_div = f"Excel tem {len(excel_rows)} linhas, reextração tem {len(comments)}."
            else:
                excel_status = "Conforme"
                excel_div = "Exportação completa bate com a reextração independente."
            out["excel_tests"].append(
                {
                    "teste": "Excel completo",
                    "projeto": label,
                    "issue": iid,
                    "resultado": excel_status,
                    "divergencia": excel_div,
                }
            )

    wrong = []
    for iid in list(issues_cti)[:8]:
        team = cti_client.request("GET", f"{cti_client.org_base}/_apis/wit/workitems/{iid}?fields=System.TeamProject&api-version=7.1")["fields"]["System.TeamProject"]
        if team != row_cti.project:
            wrong.append(("CTI", iid, team))
    for iid in list(issues_pac)[:8]:
        team = pac_client.request("GET", f"{pac_client.org_base}/_apis/wit/workitems/{iid}?fields=System.TeamProject&api-version=7.1")["fields"]["System.TeamProject"]
        if team != row_pac.project:
            wrong.append(("PAC", iid, team))

    out["filter_tests"] = {
        "cti_query_id": row_cti.id,
        "pac_query_id": row_pac.id,
        "overlap_issue_ids": sorted(set(issues_cti) & set(issues_pac)),
        "wrong_project_samples": wrong,
        "cti_indicators_vs_result": {
            "issues_found": row_cti.issues_found,
            "stored_issues": len(issues_cti),
            "comments_analyzed": row_cti.comments_analyzed,
            "sum_comment_count": sum(int(i.get("comment_count") or 0) for i in cti_ok),
            "comments_matched": row_cti.comments_matched,
            "sum_person_comments": sum(int(i.get("person_comment_count") or 0) for i in cti_ok),
        },
        "pac_indicators_vs_result": {
            "issues_found": row_pac.issues_found,
            "stored_issues": len(issues_pac),
            "comments_analyzed": row_pac.comments_analyzed,
            "sum_comment_count": sum(int(i.get("comment_count") or 0) for i in pac_ok),
            "comments_matched": row_pac.comments_matched,
            "sum_person_comments": sum(int(i.get("person_comment_count") or 0) for i in pac_ok),
        },
    }

    all_tests = out["count_tests"] + out["authorship_tests"] + out["false_negative_tests"] + out["integrity_tests"] + out["excel_tests"]
    divergent = [t for t in all_tests if t["resultado"] == "Divergente"]
    inconclusive = [t for t in all_tests if t["resultado"] == "Inconclusivo"]
    out["summary"] = {
        "issues_verificadas": len({t["issue"] for t in all_tests}),
        "testes_executados": len(all_tests),
        "divergencias": len(divergent),
        "inconclusivos": len(inconclusive),
        "confiavel_nos_casos_analisados": len(divergent) == 0 and len(inconclusive) == 0,
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))
    cti_client.close()
    pac_client.close()
    db.close()


if __name__ == "__main__":
    main()
