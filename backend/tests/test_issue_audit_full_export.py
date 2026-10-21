from io import BytesIO

from openpyxl import load_workbook

from app.services.issue_audit.comment_storage import full_comments_exportable
from app.services.issue_audit.excel_export import build_audit_workbook
from app.services.issue_audit.matching import PersonTarget
from app.services.issue_audit.runner import execute_audit
from tests.test_issue_audit import FakeCollection, _update


def _result_with_mixed_comments() -> dict:
    marcia = "Marcia Regina De Carvalho Buhr"
    fake = FakeCollection(
        {
            10: [
                _update("1", marcia, "a1", author_id="marcia-id"),
                _update("2", "Outra", "b1", author_id="x"),
                _update("3", marcia, "a2", author_id="marcia-id"),
            ],
            11: [_update("4", "Outra", "only other", author_id="x")],
            12: [],
        }
    )
    return execute_audit(
        fake,
        project="Projeto",
        work_item_type="Issue",
        period_criterion="created",
        year=2025,
        comment_scope="all",
        person=PersonTarget(marcia, identity_id="marcia-id"),
    )


def test_legacy_result_not_exportable() -> None:
    legacy = {
        "comments_analyzed": 3,
        "issues": [
            {"id": 1, "comment_count": 3, "comments": [{"text": "one"}], "error": None},
        ],
    }
    ok, message = full_comments_exportable(legacy)
    assert ok is False
    assert message


def test_storage_version_without_full_lists_not_exportable() -> None:
    legacy = {
        "comments_storage_version": 1,
        "comments_analyzed": 10,
        "issues": [{"id": 1, "comment_count": 10, "comments": [{"text": "a"}], "all_comments": [{"text": "a"}]}],
    }
    ok, message = full_comments_exportable(legacy)
    assert ok is False
    assert message


def test_full_export_sheet_rows_and_flags() -> None:
    result = _result_with_mixed_comments()
    query = {
        "organization": "CTI",
        "project": "Projeto",
        "work_item_type": "Issue",
        "period_criterion": "created",
        "year": 2025,
        "comment_scope": "all",
        "person_name": "Marcia Regina De Carvalho Buhr",
        "person_identity_id": "marcia-id",
    }
    buffer, _split = build_audit_workbook(
        query,
        result,
        base_url="https://azure-devops.tjmt.jus.br",
        include_all_discussions=True,
    )
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    assert "Discussões completas" in workbook.sheetnames
    assert "Comentários encontrados" not in workbook.sheetnames
    sheet = workbook["Discussões completas"]
    assert sheet.max_row == 5  # header + 3 + 1 comentários (Issue 12 sem linhas)
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    issue10 = [row for row in rows if row[0] == 10]
    assert len(issue10) == 3
    assert all(row[7] == "Sim" for row in issue10)
    assert sum(1 for row in issue10 if row[8] == "Sim") == 2
    issue11 = [row for row in rows if row[0] == 11]
    assert len(issue11) == 1
    assert issue11[0][7] == "Não"
    assert issue11[0][8] == "Não"


def test_traditional_export_unchanged_sheet_name() -> None:
    result = _result_with_mixed_comments()
    query = {
        "organization": "CTI",
        "project": "Projeto",
        "work_item_type": "Issue",
        "period_criterion": "created",
        "year": 2025,
        "comment_scope": "all",
        "person_name": "Marcia Regina De Carvalho Buhr",
        "person_identity_id": "marcia-id",
    }
    buffer, _split = build_audit_workbook(query, result, base_url="https://example.test", include_all_discussions=False)
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    assert workbook.sheetnames[:3] == ["Resumo das Issues", "Comentários encontrados", "Parâmetros da Consulta"]
    comments = workbook["Comentários encontrados"]
    assert comments.max_row == 3  # header + 2 da Marcia
