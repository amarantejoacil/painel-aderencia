from datetime import date
from io import BytesIO

import httpx
from openpyxl import load_workbook

from app.services.azure_devops_errors import AzureDevOpsConnectionError
from app.services.issue_audit.client import AzureCollectionClient, work_item_url
from app.services.issue_audit.excel_export import build_audit_workbook, html_to_plain_text, safe_excel_text
from app.services.issue_audit.matching import PersonTarget, comment_from_update, person_matches
from app.services.issue_audit.runner import execute_audit


def _update(update_id: str, author: str, text: str, *, author_id: str | None = None, when: str = "2025-06-01T15:00:00Z") -> dict:
    revised = {"displayName": author, "uniqueName": f"{author_id or 'user'}@tjmt.jus.br"}
    if author_id:
        revised["id"] = author_id
    return {
        "id": update_id,
        "revisedDate": when,
        "revisedBy": revised,
        "fields": {"System.History": {"newValue": text}} if text is not None else {"System.State": {"newValue": "Active"}},
    }


class FakeCollection:
    field_catalog: list[dict] = []

    def __init__(self, updates: dict[int, list[dict]], failures: set[int] | None = None) -> None:
        self.base_url = "https://azure-devops.tjmt.jus.br"
        self.organization = "CTI"
        self.updates = updates
        self.failures = failures or set()
        self.wiql = None

    def ensure_project(self, project: str) -> None:
        self.project = project

    def query_work_item_ids(self, **kwargs):
        self.wiql = kwargs
        return list(self.updates), []

    def fetch_work_items(self, ids: list[int], *, field_refs=None):
        del field_refs
        items = []
        for item_id in ids:
            items.append(
                {
                    "id": item_id,
                    "fields": {
                        "System.Title": f"Issue {item_id}",
                        "System.WorkItemType": "Issue",
                        "System.State": "Active",
                        "System.AssignedTo": {"displayName": "Outra Pessoa"},
                        "System.CreatedDate": "2025-02-01T12:00:00Z",
                        "System.ChangedDate": "2025-03-01T12:00:00Z",
                    },
                }
            )
        return items, []

    def fetch_comment_updates(self, work_item_id: int):
        if work_item_id in self.failures:
            raise AzureDevOpsConnectionError("Falha ao consultar comentários.")
        return self.updates[work_item_id]


def test_history_without_text_is_not_a_comment() -> None:
    assert comment_from_update(_update("1", "Alguém", None)) is None
    comment = comment_from_update(_update("2", "Marcia Regina De Carvalho Buhr", "Texto original"))
    assert comment is not None
    assert comment["text"] == "Texto original"
    assert comment["is_deleted"] is None


def test_mention_in_text_is_not_authorship() -> None:
    target = PersonTarget("Marcia Regina De Carvalho Buhr", identity_id="marcia-id")
    mention = {
        "author": "Outra Pessoa",
        "author_id": "outra-id",
        "author_unique_name": "outra@tjmt.jus.br",
        "text": "Marcia Regina De Carvalho Buhr deveria ver isso",
    }
    own = {
        "author": "Marcia Regina De Carvalho Buhr",
        "author_id": "marcia-id",
        "author_unique_name": "marcia@tjmt.jus.br",
        "text": "registro",
    }
    assert person_matches(mention, target) is False
    assert person_matches(own, target) is True


def test_audit_counts_person_comments_and_keeps_issues_without_them() -> None:
    marcia = "Marcia Regina De Carvalho Buhr"
    fake = FakeCollection(
        {
            200: [
                _update("1", marcia, "primeiro", author_id="marcia-id", when="2025-04-01T14:00:00Z"),
                _update("2", "Outra Pessoa", f"citei {marcia}", author_id="outra-id"),
                _update("3", marcia, "segundo", author_id="marcia-id", when="2025-05-02T14:00:00Z"),
                _update("4", marcia, None, author_id="marcia-id"),
            ],
            201: [_update("5", "Outra Pessoa", "sem a pessoa", author_id="outra-id")],
            202: [],
        },
        failures={203},
    )
    fake.updates[203] = []
    result = execute_audit(
        fake,
        project="Gestão Gabinete Coordenadoria",
        work_item_type="Issue",
        period_criterion="created",
        year=2025,
        comment_scope="all",
        person=PersonTarget(marcia),
    )
    by_id = {issue["id"]: issue for issue in result["issues"]}
    assert result["status"] == "partial"
    assert result["issues_found"] == 4
    assert by_id[200]["comment_count"] == 3
    assert by_id[200]["person_comment_count"] == 2
    assert [item["text"] for item in by_id[200]["comments"]] == ["primeiro", "segundo"]
    assert len(by_id[200]["all_comments"]) == 3
    assert result.get("comments_storage_version") == 1
    assert by_id[201]["person_commented"] is False
    assert by_id[202]["comment_count"] == 0
    assert by_id[202]["person_commented"] is False
    assert by_id[203]["error"]
    assert result["failure_count"] == 1
    assert fake.wiql["date_field"] == "System.CreatedDate"
    assert fake.wiql["start"] == date(2025, 1, 1)
    assert fake.wiql["end"] == date(2026, 1, 1)
    assert result["person"]["identity_id"] == "marcia-id"
    assert by_id[200]["url"] == work_item_url(fake.base_url, "CTI", "Gestão Gabinete Coordenadoria", 200)


def test_same_display_name_with_two_ids_is_ambiguous() -> None:
    name = "Marcia Regina De Carvalho Buhr"
    fake = FakeCollection(
        {
            10: [_update("1", name, "da conta A", author_id="a")],
            11: [_update("2", name, "da conta B", author_id="b")],
        }
    )
    result = execute_audit(
        fake,
        project="Projeto",
        work_item_type="Issue",
        period_criterion="changed",
        year=2025,
        comment_scope="all",
        person=PersonTarget(name),
    )
    assert result["status"] == "ambiguous"
    assert len(result["candidates"]) == 2
    assert result["comments_matched"] == 2


def test_comment_outside_year_is_excluded_in_cuiaba() -> None:
    name = "Marcia Regina De Carvalho Buhr"
    fake = FakeCollection(
        {
            10: [
                _update("1", name, "ainda 2024 em Cuiabá", author_id="marcia-id", when="2025-01-01T02:00:00Z"),
                _update("2", name, "dentro de 2025", author_id="marcia-id", when="2025-01-01T05:00:00Z"),
            ]
        }
    )
    result = execute_audit(
        fake,
        project="Projeto",
        work_item_type="Issue",
        period_criterion="created",
        year=2025,
        comment_scope="period",
        person=PersonTarget(name, identity_id="marcia-id"),
    )
    issue = result["issues"][0]
    assert issue["comment_count"] == 1
    assert issue["comments"][0]["text"] == "dentro de 2025"


def test_update_pagination_and_retry() -> None:
    calls = {"updates": 0, "wiql": 0}

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/updates"):
            calls["updates"] += 1
            skip = int(request.url.params.get("$skip", "0"))
            if skip == 0:
                page = [{"id": index, "fields": {}} for index in range(200)]
            else:
                page = [{"id": 200, "fields": {"System.History": {"newValue": "último"}}}]
            return httpx.Response(200, json={"count": len(page), "value": page})
        if request.url.path.endswith("/wiql"):
            calls["wiql"] += 1
            if calls["wiql"] == 1:
                return httpx.Response(429, headers={"Retry-After": "0"})
            return httpx.Response(200, json={"workItems": [{"id": 7}]})
        return httpx.Response(404)

    client = AzureCollectionClient(
        base_url="https://azure-devops.tjmt.jus.br",
        organization="CTI",
        pat="pat",
        sleeper=lambda _seconds: None,
        http_client=httpx.Client(transport=httpx.MockTransport(handler)),
    )
    updates = client.fetch_comment_updates(200)
    assert len(updates) == 201
    assert calls["updates"] == 2
    ids, notes = client.query_work_item_ids(
        project="Gestão",
        work_item_type="Issue",
        date_field="System.CreatedDate",
        start=date(2025, 1, 1),
        end=date(2026, 1, 1),
    )
    assert ids == [7]
    assert notes == []
    client.close()


def test_html_to_plain_text_normalizes_azure_line_breaks() -> None:
    assert "linha um\n\nlinha dois" in html_to_plain_text("linha um_x000d__x000d_linha dois")


def test_html_to_plain_text_strips_tags_and_decodes_entities() -> None:
    raw = (
        '<div>Lançar no sistema de diárias</div>'
        '<a href="#" data-vss-mention="x">@Luisa Nunes Guimaraes</a>'
        "&nbsp;encaminhar ao DSA"
    )
    plain = html_to_plain_text(raw)
    assert "<div>" not in plain
    assert "<a " not in plain
    assert "Lançar no sistema de diárias" in plain
    assert "@Luisa Nunes Guimaraes" in plain
    assert " encaminhar ao DSA" in plain or "encaminhar ao DSA" in plain


def test_excel_has_three_sheets_formula_guard_and_split() -> None:
    long_text = "linha\n" + ("A" * 40000)
    result = {
        "status": "partial",
        "issues_found": 2,
        "comments_analyzed": 2,
        "comments_matched": 1,
        "failure_count": 1,
        "limitations": ["falha parcial"],
        "error_message": None,
        "person": {"identity_id": "marcia-id"},
        "issues": [
            {
                "id": 200,
                "title": "=cmd|calc",
                "description": "<div>Texto da issue</div>",
                "tags": "DSI",
                "bucket": "Concluído",
                "areas_internas": "DSI",
                "areas_externa": "",
                "assigned_to": "Alguém",
                "state": "Active",
                "created_at": "2025-02-01T15:00:00Z",
                "comment_count": 1,
                "person_comment_count": 1,
                "person_commented": True,
                "url": "https://azure-devops.tjmt.jus.br/CTI/Projeto/_workitems/edit/200",
                "comments": [
                    {
                        "author": "Marcia Regina De Carvalho Buhr",
                        "created_at": "2025-04-01T18:30:00Z",
                        "text": '<div>linha</div><br/>' + ("A" * 39990),
                    }
                ],
            },
            {
                "id": 201,
                "title": "Sem comentário",
                "assigned_to": "Outra",
                "state": "Closed",
                "created_at": "2025-02-02T15:00:00Z",
                "comment_count": 0,
                "person_comment_count": 0,
                "person_commented": False,
                "url": "https://evil.example/issue/201",
                "comments": [],
            },
        ],
    }
    query = {
        "organization": "CTI",
        "project": "Gestão Gabinete Coordenadoria",
        "work_item_type": "Issue",
        "period_criterion": "created",
        "year": 2025,
        "comment_scope": "all",
        "person_name": "Marcia Regina De Carvalho Buhr",
        "person_identity_id": "marcia-id",
    }
    buffer, split = build_audit_workbook(query, result, base_url="https://azure-devops.tjmt.jus.br")
    assert split is True
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    assert workbook.sheetnames == ["Resumo das Issues", "Comentários encontrados", "Parâmetros da Consulta"]
    summary = workbook["Resumo das Issues"]
    assert summary.freeze_panes == "A2"
    assert summary.auto_filter.ref.startswith("A1:")
    assert summary["A2"].hyperlink.target.endswith("/200")
    assert summary["B2"].value.startswith("'=")
    assert summary["C2"].value == "Texto da issue"
    assert summary["D2"].value == "DSI"
    assert summary["E2"].value == "Concluído"
    assert summary["F2"].value == "DSI"
    assert summary["A3"].hyperlink is None
    assert summary["M2"].value == "Sim"
    assert summary["M3"].value == "Não"
    comments = workbook["Comentários encontrados"]
    assert comments["G2"].value.startswith("linha")
    assert "<div>" not in str(comments["G2"].value)
    assert comments.max_row == 3
    assert "continuação" in " ".join(str(cell.value) for cell in workbook["Parâmetros da Consulta"]["B"] if cell.value).casefold()
    assert safe_excel_text("=1+1").startswith("'")
