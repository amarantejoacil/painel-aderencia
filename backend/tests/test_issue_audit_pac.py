from datetime import date
from io import BytesIO

import httpx
from openpyxl import load_workbook

from app.services.issue_audit.client import AzureCollectionClient
from app.services.issue_audit.excel_export import build_audit_workbook
from app.services.issue_audit.matching import PersonTarget, comment_from_api, person_matches
from app.services.issue_audit.profiles import NUCLEO_PAC, PAC_WORK_ITEM_FIELDS
from app.services.issue_audit.runner import execute_audit

MARCIA = "Marcia Regina De Carvalho Buhr"
NUCLEO_ID = "9982c2e2-276e-480d-b049-a29c7d1a80e5"
CTI_ID = "1e66179c-7ef1-40b6-8533-97418539f742"
TITLE = (
    "Aquisição de Equipamentos de TIC - Mini PC Mac studio Tablet e Notebook  "
    "(Desembargadores) impressora e scanner portátil"
)


def _api_comment(comment_id: int, author: str, author_id: str, text: str, unique: str = "PJMT\\1") -> dict:
    return {
        "id": comment_id,
        "text": text,
        "createdDate": f"2026-07-{(comment_id % 27) + 1:02d}T18:00:00Z",
        "createdBy": {"displayName": author, "id": author_id, "uniqueName": unique},
    }


def _issue_26689_comments() -> list[dict]:
    comments = [
        _api_comment(index, "FABIO BARROS DE MIRALLA SANTOS", "fabio-id", f"comentario {index}")
        for index in range(1, 56)
    ]
    comments.append(
        _api_comment(56, "DANIELLA CHRISTINA BARROS DE CARVALHO", "dani-id", f"{MARCIA} citada no texto")
    )
    comments.append(_api_comment(57, "JOACIL AMARANTE DE PAULA JUNIOR", "joacil-id", "abertura"))
    comments.extend(
        _api_comment(100 + index, MARCIA, NUCLEO_ID, f"registro {index}", "PJMT\\55940")
        for index in range(1, 4)
    )
    comments.append({"id": 999, "text": "   ", "createdBy": {"displayName": MARCIA, "id": NUCLEO_ID}})
    return comments


class PacCollection:
    def __init__(self, pages: list[list[dict]]) -> None:
        self.base_url = "https://azure-devops.tjmt.jus.br"
        self.organization = "Nucleo"
        self.pages = pages
        self.field_refs = None
        self.wiql = None
        self.calls = 0

    def ensure_project(self, project: str) -> None:
        self.project = project

    def query_work_item_ids(self, **kwargs):
        self.wiql = kwargs
        return [26689], []

    def fetch_work_items(self, ids: list[int], *, field_refs=None):
        self.field_refs = list(field_refs or [])
        return [
            {
                "id": 26689,
                "fields": {
                    "System.Title": TITLE,
                    "System.WorkItemType": "Issue",
                    "System.TeamProject": NUCLEO_PAC.project,
                    "System.State": "Em andamento",
                    "System.AssignedTo": {"displayName": "FABIO BARROS DE MIRALLA SANTOS"},
                    "System.CreatedDate": "2026-02-06T17:29:42.737Z",
                    "System.ChangedDate": "2026-09-29T16:45:01.237Z",
                },
            }
            for _item_id in ids
        ], []

    def fetch_work_item_comments(self, project: str, work_item_id: int):
        del project, work_item_id
        self.calls += 1
        return [item for page in self.pages for item in page]


def test_empty_api_comment_is_ignored_and_mention_is_not_authorship() -> None:
    assert comment_from_api({"id": 1, "text": "  ", "createdBy": {"displayName": MARCIA, "id": NUCLEO_ID}}) is None
    mention = comment_from_api(_api_comment(2, "Outra Pessoa", "outra", f"encaminhar para {MARCIA}"))
    assert mention is not None
    assert person_matches(mention, PersonTarget(MARCIA, identity_id=NUCLEO_ID)) is False


def test_cti_identity_does_not_match_nucleo_comments() -> None:
    own = comment_from_api(_api_comment(3, MARCIA, NUCLEO_ID, "meu registro", "PJMT\\55940"))
    assert own is not None
    assert person_matches(own, PersonTarget(MARCIA, identity_id=CTI_ID, unique_name="55940@tjmt.jus.br")) is False
    assert person_matches(own, PersonTarget(MARCIA, identity_id=NUCLEO_ID, unique_name="55940@tjmt.jus.br")) is True


def test_pac_profile_reads_sixty_comments_without_cti_fields() -> None:
    fake = PacCollection([_issue_26689_comments()])
    result = execute_audit(
        fake,
        project=NUCLEO_PAC.project,
        work_item_type="Issue",
        period_criterion="created",
        year=2026,
        comment_scope="all",
        person=PersonTarget(MARCIA),
        comment_strategy="comments_api",
        planning_fields=False,
        profile_id=NUCLEO_PAC.id,
        state="Em andamento",
    )
    issue = result["issues"][0]
    assert result["comment_source"] == "comments_api"
    assert result["profile_id"] == "nucleo-pac"
    assert issue["title"] == TITLE
    assert issue["assigned_to"] == "FABIO BARROS DE MIRALLA SANTOS"
    assert issue["state"] == "Em andamento"
    assert issue["comment_count"] == 60
    assert issue["person_comment_count"] == 3
    assert issue["person_commented"] is True
    assert len(issue["all_comments"]) == 60
    assert result["person"]["identity_id"] == NUCLEO_ID
    assert result["comments_matched"] == 3
    assert fake.field_refs == PAC_WORK_ITEM_FIELDS
    assert "Custom.Bucket" not in fake.field_refs
    assert fake.wiql["state"] == "Em andamento"
    assert fake.wiql["start"] == date(2026, 1, 1)
    assert "System.History" in result["limitations"][0] or "API de comentários" in result["limitations"][0]
    assert "System.History" not in result["limitations"][0]


def test_custom_period_replaces_calendar_year() -> None:
    fake = PacCollection([_issue_26689_comments()])
    execute_audit(
        fake,
        project=NUCLEO_PAC.project,
        work_item_type="Issue",
        period_criterion="changed",
        year=2025,
        comment_scope="all",
        person=PersonTarget(MARCIA, identity_id=NUCLEO_ID),
        comment_strategy="comments_api",
        planning_fields=False,
        period_start=date(2026, 2, 1),
        period_end=date(2026, 2, 28),
    )
    assert fake.wiql["start"] == date(2026, 2, 1)
    assert fake.wiql["end"] == date(2026, 3, 1)
    assert fake.wiql["date_field"] == "System.ChangedDate"


def test_pac_excel_omits_cti_columns_and_keeps_yes_no() -> None:
    fake = PacCollection([_issue_26689_comments()])
    result = execute_audit(
        fake,
        project=NUCLEO_PAC.project,
        work_item_type="Issue",
        period_criterion="created",
        year=2026,
        comment_scope="all",
        person=PersonTarget(MARCIA, identity_id=NUCLEO_ID),
        comment_strategy="comments_api",
        planning_fields=False,
        profile_id="nucleo-pac",
    )
    query = {
        "organization": "Nucleo",
        "project": NUCLEO_PAC.project,
        "work_item_type": "Issue",
        "period_criterion": "created",
        "year": 2026,
        "comment_scope": "all",
        "person_name": MARCIA,
        "person_identity_id": NUCLEO_ID,
        "profile_id": "nucleo-pac",
    }
    buffer, _split = build_audit_workbook(
        query,
        result,
        base_url="https://azure-devops.tjmt.jus.br",
        include_all_discussions=False,
    )
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    summary = workbook["Resumo das Issues"]
    headers = [cell.value for cell in summary[1]]
    assert "Bucket" not in headers
    assert "Áreas Internas" not in headers
    assert headers[7] == "Pessoa pesquisada comentou?"
    assert summary["H2"].value == "Sim"
    assert summary["A2"].hyperlink.target.endswith("/26689")
    person_sheet = workbook["Comentários encontrados"]
    person_headers = [cell.value for cell in person_sheet[1]]
    assert person_headers[7] == "Pessoa pesquisada comentou nesta Issue?"
    assert person_headers[8] == "Este comentário é da pessoa pesquisada?"
    assert person_sheet.max_row == 4
    assert all(row[7] == "Sim" and row[8] == "Sim" for row in person_sheet.iter_rows(min_row=2, values_only=True))

    full_buffer, _split = build_audit_workbook(
        query,
        result,
        base_url="https://azure-devops.tjmt.jus.br",
        include_all_discussions=True,
    )
    full = load_workbook(BytesIO(full_buffer.getvalue()))
    discussions = full["Discussões completas"]
    rows = list(discussions.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 60
    assert all(row[7] == "Sim" for row in rows)
    assert sum(1 for row in rows if row[8] == "Sim") == 3
    parameters = " ".join(
        str(value)
        for row in full["Parâmetros da Consulta"].iter_rows(values_only=True)
        for value in row
        if value is not None
    )
    assert "Nucleo" in parameters
    assert "7.1-preview.4" in parameters
    assert "Discussões completas" in parameters


def test_comments_api_follows_continuation_token() -> None:
    calls: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/comments"):
            calls.append(request.url.params.get("continuationToken", ""))
            if not request.url.params.get("continuationToken"):
                return httpx.Response(
                    200,
                    json={
                        "totalCount": 3,
                        "comments": [_api_comment(1, MARCIA, NUCLEO_ID, "um")],
                        "continuationToken": "proxima",
                    },
                )
            return httpx.Response(
                200,
                json={"totalCount": 3, "comments": [_api_comment(2, "Outra", "x", "dois"), _api_comment(3, MARCIA, NUCLEO_ID, "tres")]},
            )
        return httpx.Response(404)

    client = AzureCollectionClient(
        base_url="https://azure-devops.tjmt.jus.br",
        organization="Nucleo",
        pat="pat",
        sleeper=lambda _seconds: None,
        http_client=httpx.Client(transport=httpx.MockTransport(handler)),
    )
    comments = client.fetch_work_item_comments("Núcleo de Aquisições - PAC", 26689)
    assert [item["id"] for item in comments] == [1, 2, 3]
    assert calls == ["", "proxima"]
    assert "/Nucleo/" in str(client.project_base("Núcleo de Aquisições - PAC"))
    client.close()
