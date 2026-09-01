from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.models import Activity, Collaborator
from app.services.activities_export import (
    assignee_label,
    build_activities_workbook,
    build_export_filename,
)
from app.services.azure_state_labels import translate_azure_state


def _activity(**overrides) -> Activity:
    data = dict(
        task_id="12345",
        title="Reunião de acompanhamento",
        work_item_type="Task",
        azure_assignee="João da Silva",
        collaborator_id=1,
        work_date=date(2026, 8, 3),
        estimated_hours=None,
        completed_hours=Decimal("2.00"),
        state="Closed",
        project=None,
        activity_category="Desenvolvimento",
        import_id=1,
    )
    data.update(overrides)
    activity = Activity(**data)
    activity.collaborator = Collaborator(
        id=1,
        name="João da Silva",
        azure_name="João da Silva",
        start_date=date(2026, 8, 1),
        end_date=None,
        daily_hours=Decimal("8.00"),
        active=True,
    )
    return activity


def test_translate_azure_state_closed() -> None:
    assert translate_azure_state("Closed") == "Concluído"
    assert translate_azure_state("closed") == "Concluído"
    assert translate_azure_state("Custom") == "Custom"


def test_build_export_filename() -> None:
    collaborator = Collaborator(
        id=5,
        name="Joacil Amarante de Paula Junior",
        azure_name="JOACIL",
        start_date=date(2026, 8, 3),
        end_date=None,
        daily_hours=Decimal("8.00"),
        active=True,
    )
    assert build_export_filename(2026, 8) == "atividades_agosto_2026.xlsx"
    assert (
        build_export_filename(2026, 8, collaborator)
        == "atividades_joacil_amarante_de_paula_junior_agosto_2026.xlsx"
    )


def test_build_activities_workbook_structure() -> None:
    activities = [
        _activity(task_id="12345", completed_hours=Decimal("2"), state="Closed"),
        _activity(
            task_id="12346",
            title="Análise de requisitos",
            completed_hours=Decimal("6"),
            state="Closed",
        ),
    ]
    buffer = build_activities_workbook(activities)
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    sheet = workbook["Atividades"]

    assert sheet.cell(1, 1).value == "ID"
    assert sheet.cell(1, 4).value == "Atividade"
    assert sheet.cell(1, 7).value == "Horas Executadas"
    assert sheet.cell(2, 1).value == "12345"
    assert sheet.cell(2, 4).value == "Desenvolvimento"
    assert sheet.cell(2, 6).value == "Concluído"
    assert sheet.cell(2, 7).value == 2.0
    assert sheet.cell(3, 7).value == 6.0
    assert assignee_label(activities[0]) == "João da Silva"
    assert sheet.auto_filter.ref == "A1:G3"
