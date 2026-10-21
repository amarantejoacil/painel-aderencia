from io import BytesIO

from openpyxl import load_workbook

from app.models import IssueAuditQuery, IssueAuditValidation
from app.services.issue_audit.excel_export import build_audit_workbook
from app.services.issue_audit.validation import (
    build_stratified_sample,
    compute_metrics,
    extraction_diagnosis,
    save_manual_review,
    validate_coverage,
)


class FakeClient:
    organization = "CTI"
    base_url = "https://azure-devops.tjmt.jus.br"

    def __init__(self, wiql_ids: list[int]) -> None:
        self._wiql_ids = wiql_ids

    def query_work_item_ids(self, **kwargs):
        del kwargs
        return list(self._wiql_ids), []

    def fetch_work_items(self, ids: list[int], *, field_refs=None):
        del field_refs
        items = [
            {
                "id": item_id,
                "fields": {
                    "System.WorkItemType": "Issue",
                    "System.TeamProject": "Gestão Gabinete Coordenadoria",
                },
            }
            for item_id in ids
        ]
        return items, []


def test_extraction_diagnosis_has_risks() -> None:
    diag = extraction_diagnosis()
    assert diag["source"] == "system_history"
    assert diag["limitation"]
    assert len(diag["risks"]) >= 3


def test_validate_coverage_detects_missing_and_extra() -> None:
    row = IssueAuditQuery(
        id=3,
        organization="CTI",
        project="Gestão Gabinete Coordenadoria",
        work_item_type="Issue",
        period_criterion="created",
        year=2025,
        comment_scope="all",
        person_name="Marcia Regina De Carvalho Buhr",
        status="complete",
        issues_found=3,
        result={"issues": [{"id": 1}, {"id": 2}, {"id": 99}]},
    )
    client = FakeClient([1, 2, 3])
    report = validate_coverage(client, row)
    assert report["consistent"] is False
    assert report["missing_count"] == 1
    assert 3 in report["missing_in_stored"]
    assert report["extra_count"] == 1
    assert 99 in report["extra_in_stored"]


def test_build_stratified_sample_covers_strata() -> None:
    issues = [
        {"id": 1, "person_commented": True, "comment_count": 5, "created_at": "2025-01-02"},
        {"id": 2, "person_commented": False, "comment_count": 0, "created_at": "2025-04-02"},
        {"id": 3, "person_commented": False, "comment_count": 12, "created_at": "2025-07-02"},
        {"id": 4, "person_commented": False, "comment_count": 1, "created_at": "2025-10-02"},
    ]
    row = IssueAuditQuery(
        id=1,
        organization="CTI",
        project="P",
        work_item_type="Issue",
        period_criterion="created",
        year=2025,
        comment_scope="all",
        person_name="X",
        status="complete",
        issues_found=4,
    )
    sample = build_stratified_sample(issues, row)
    assert set(sample) >= {1, 2, 3}


class _FakeSession:
    def commit(self) -> None:
        return None

    def refresh(self, _obj) -> None:
        return None


def test_manual_review_updates_validation_payload_only() -> None:
    query = IssueAuditQuery(
        id=3,
        organization="CTI",
        project="P",
        work_item_type="Issue",
        period_criterion="created",
        year=2025,
        comment_scope="all",
        person_name="Marcia",
        status="complete",
        issues_found=1,
        comments_matched=1,
        result={"issues": [{"id": 200, "person_comment_count": 1}]},
    )
    validation = IssueAuditValidation(
        id=1,
        query_id=3,
        created_at=__import__("datetime").datetime.utcnow(),
        payload={
            "sample": {"issue_ids": [200], "items": []},
            "manual_reviews": [],
            "metrics": compute_metrics({"consistent": True}, [], query),
        },
    )
    save_manual_review(
        _FakeSession(),
        validation,
        issue_id=200,
        status="divergencia",
        real_comment_count=2,
        real_person_comment_count=1,
        divergence_type="contagem",
        notes="teste",
        reviewer="auditor",
    )
    assert validation.payload["manual_reviews"][0]["status"] == "divergencia"
    assert query.result["issues"][0]["person_comment_count"] == 1


def test_excel_includes_validacao_sheet() -> None:
    validation = {
        "validated_at": "2025-10-21T12:00:00+00:00",
        "methodology": {"limitation": "System.History"},
        "metrics": {"disclaimer": "Amostral"},
        "sample": {
            "items": [
                {
                    "issue_id": 200,
                    "url": "https://azure-devops.tjmt.jus.br/CTI/P/_workitems/edit/200",
                    "original": {"comment_count": 3, "person_comment_count": 1},
                    "reextracted": {"comment_count": 3, "person_comment_count": 1},
                    "auto_status": "confere",
                    "divergence": None,
                }
            ]
        },
        "manual_reviews": [],
    }
    buffer, _ = build_audit_workbook(
        {"organization": "CTI", "project": "P", "work_item_type": "Issue", "period_criterion": "created", "year": 2025, "comment_scope": "all", "person_name": "M"},
        {"issues": [{"id": 200, "title": "T", "assigned_to": "A", "state": "Active", "created_at": "2025-01-01", "comment_count": 3, "person_comment_count": 1, "person_commented": True, "url": "https://azure-devops.tjmt.jus.br/CTI/P/_workitems/edit/200", "comments": []}]},
        base_url="https://azure-devops.tjmt.jus.br",
        validation=validation,
    )
    workbook = load_workbook(BytesIO(buffer.getvalue()))
    assert "Validação" in workbook.sheetnames
