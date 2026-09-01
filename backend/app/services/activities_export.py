from __future__ import annotations

import re
import unicodedata
from calendar import monthrange
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models import Activity, Collaborator, ImportBatch
from app.services.azure_state_labels import translate_azure_state

HEADERS = [
    "ID",
    "Data da Atividade",
    "Título",
    "Atividade",
    "Atribuído para",
    "Status da Atividade",
    "Horas Executadas",
]

HEADER_FILL = PatternFill("solid", fgColor="1F6B59")
HEADER_FONT = Font(bold=True, color="FFFFFF")
COLUMN_WIDTHS = {
    "A": 12,
    "B": 16,
    "C": 48,
    "D": 28,
    "E": 28,
    "F": 22,
    "G": 16,
}


def _slug_name(name: str) -> str:
    normalized = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "_", normalized.lower()).strip("_")
    return slug or "colaborador"


def _month_name_pt(month: int) -> str:
    names = [
        "",
        "janeiro",
        "fevereiro",
        "marco",
        "abril",
        "maio",
        "junho",
        "julho",
        "agosto",
        "setembro",
        "outubro",
        "novembro",
        "dezembro",
    ]
    return names[month]


def fetch_activities(
    db: Session,
    year: int,
    month: int,
    collaborator_id: int | None = None,
) -> list[Activity]:
    latest = db.scalars(select(ImportBatch).order_by(ImportBatch.imported_at.desc())).first()
    if not latest:
        return []

    first = date(year, month, 1)
    last = date(year, month, monthrange(year, month)[1])
    stmt = (
        select(Activity)
        .options(joinedload(Activity.collaborator))
        .where(
            Activity.import_id == latest.id,
            Activity.work_date >= first,
            Activity.work_date <= last,
        )
        .order_by(Activity.work_date, Activity.task_id)
    )
    if collaborator_id is not None:
        stmt = stmt.where(Activity.collaborator_id == collaborator_id)
    return list(db.scalars(stmt).unique().all())


def assignee_label(activity: Activity) -> str:
    if activity.collaborator is not None:
        return activity.collaborator.name
    return activity.azure_assignee


def build_export_filename(
    year: int,
    month: int,
    collaborator: Collaborator | None = None,
) -> str:
    month_label = _month_name_pt(month)
    if collaborator is not None:
        return f"atividades_{_slug_name(collaborator.name)}_{month_label}_{year}.xlsx"
    return f"atividades_{month_label}_{year}.xlsx"


def _hours_value(value: Decimal | None) -> float | None:
    if value is None:
        return None
    return float(value)


def build_activities_workbook(activities: list[Activity]) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Atividades"

    for col_index, header in enumerate(HEADERS, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, activity in enumerate(activities, start=2):
        worksheet.cell(row=row_index, column=1, value=activity.task_id)

        date_cell = worksheet.cell(row=row_index, column=2, value=activity.work_date)
        date_cell.number_format = "DD/MM/YYYY"
        date_cell.alignment = Alignment(horizontal="center")

        title_cell = worksheet.cell(row=row_index, column=3, value=activity.title)
        title_cell.alignment = Alignment(wrap_text=True, vertical="top")

        category_cell = worksheet.cell(row=row_index, column=4, value=activity.activity_category)
        category_cell.alignment = Alignment(wrap_text=True, vertical="top")

        worksheet.cell(row=row_index, column=5, value=assignee_label(activity))
        worksheet.cell(row=row_index, column=6, value=translate_azure_state(activity.state))

        hours_cell = worksheet.cell(row=row_index, column=7, value=_hours_value(activity.completed_hours))
        if hours_cell.value is not None:
            hours_cell.number_format = "0.00"
            hours_cell.alignment = Alignment(horizontal="right")
        else:
            hours_cell.alignment = Alignment(horizontal="center")

    for column, width in COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width

    worksheet.auto_filter.ref = f"A1:G{max(1, len(activities) + 1)}"
    worksheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
