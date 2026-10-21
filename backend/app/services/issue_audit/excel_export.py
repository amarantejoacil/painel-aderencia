from __future__ import annotations

import html
import re
from datetime import datetime
from io import BytesIO
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.issue_audit.comment_storage import count_stored_all_comments
from app.services.issue_audit.matching import (
    DISPLAY_TZ,
    PersonTarget,
    parse_azure_datetime,
    person_matches,
    to_local,
)

EXCEL_CELL_LIMIT = 32767
CHUNK_SIZE = 32000
HEADER_FILL = PatternFill("solid", fgColor="1F6B59")
HEADER_FONT = Font(bold=True, color="FFFFFF")
YES_FILL = PatternFill("solid", fgColor="E5F2EC")
YES_FONT = Font(color="185A4A", bold=True)
NO_FILL = PatternFill("solid", fgColor="F4F4F5")
NO_FONT = Font(color="3F3F46")
WRAP = Alignment(wrap_text=True, vertical="top")

SUMMARY_HEADERS = [
    "ID",
    "Título",
    "Descrição",
    "Tags",
    "Bucket",
    "Áreas Internas",
    "Áreas Externa",
    "Atribuído a",
    "Estado",
    "Data de criação",
    "Total de comentários",
    "Comentários da pessoa pesquisada",
    "Existe comentário da pessoa?",
    "Link da Issue",
]
FULL_DISCUSSION_HEADERS = [
    "ID da Issue",
    "Título da Issue",
    "Atribuído a",
    "Data do comentário",
    "Hora do comentário",
    "Autor do comentário",
    "Conteúdo completo do comentário",
    "Pessoa pesquisada comentou nesta Issue?",
    "Este comentário é da pessoa pesquisada?",
    "Link da Issue",
]
COMMENT_HEADERS = [
    "ID da Issue",
    "Título da Issue",
    "Atribuído a",
    "Autor do comentário",
    "Data do comentário",
    "Hora do comentário",
    "Texto completo do comentário",
    "Link da Issue",
    "Parte",
]


_BR_TAG = re.compile(r"<br\s*/?>", re.IGNORECASE)
_BLOCK_END = re.compile(r"</(?:p|div|li|tr|h[1-6])\s*>", re.IGNORECASE)
_ANY_TAG = re.compile(r"<[^>]+>")


def html_to_plain_text(value: object) -> str:
    """Converte HTML do Azure DevOps em texto legível para auditoria."""
    text = "" if value is None else str(value)
    if not text.strip():
        return ""
    text = _BR_TAG.sub("\n", text)
    text = _BLOCK_END.sub("\n", text)
    text = _ANY_TAG.sub("", text)
    text = html.unescape(text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"_x000d_", "\n", text, flags=re.IGNORECASE)
    return text.strip()


def description_for_excel(value: object) -> str:
    plain = html_to_plain_text(value)
    if len(plain) <= EXCEL_CELL_LIMIT:
        return safe_excel_text(plain)
    return safe_excel_text(plain[: EXCEL_CELL_LIMIT - 1] + "…")


def safe_excel_text(value: object) -> str:
    text = "" if value is None else str(value)
    probe = text.lstrip()
    if probe[:1] in {"=", "+", "-", "@"}:
        return f"'{text}"
    return text


def text_chunks(value: str) -> list[str]:
    if len(value) <= EXCEL_CELL_LIMIT:
        return [value]
    return [value[index : index + CHUNK_SIZE] for index in range(0, len(value), CHUNK_SIZE)]


VALIDATION_HEADERS = [
    "ID da Issue",
    "Link",
    "Resultado original (comentários)",
    "Resultado original (pessoa)",
    "Reextração automática (comentários)",
    "Reextração automática (pessoa)",
    "Conferência automática",
    "Conferência manual",
    "Qtd. real comentários",
    "Qtd. real da pessoa",
    "Tipo de divergência",
    "Observações",
    "Data da validação",
    "Responsável",
]


def build_audit_workbook(
    query: dict,
    result: dict,
    *,
    base_url: str,
    validation: dict | None = None,
    include_all_discussions: bool = False,
) -> tuple[BytesIO, bool]:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo das Issues"
    exported = {**result, "base_url": base_url}
    _write_summary(summary, exported)
    split = False
    if include_all_discussions:
        discussions = workbook.create_sheet("Discussões completas")
        split = _write_full_discussions(discussions, exported, query)
    else:
        comments = workbook.create_sheet("Comentários encontrados")
        split = _write_comments(comments, exported)
    parameters = workbook.create_sheet("Parâmetros da Consulta")
    _write_parameters(
        parameters,
        query,
        exported,
        split_text=split,
        validation=validation,
        include_all_discussions=include_all_discussions,
    )
    if validation:
        sheet = workbook.create_sheet("Validação")
        _write_validation(sheet, validation, base_url=base_url)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer, split


def _write_summary(sheet: Worksheet, result: dict) -> None:
    _write_header(sheet, SUMMARY_HEADERS)
    for issue in result.get("issues") or []:
        url = str(issue.get("url") or "")
        row = [
            issue.get("id"),
            safe_excel_text(issue.get("title")),
            description_for_excel(issue.get("description")),
            safe_excel_text(issue.get("tags")),
            safe_excel_text(issue.get("bucket")),
            safe_excel_text(issue.get("areas_internas")),
            safe_excel_text(issue.get("areas_externa")),
            safe_excel_text(issue.get("assigned_to")),
            safe_excel_text(issue.get("state")),
            _format_stamp(issue.get("created_at"), with_time=False),
            issue.get("comment_count") if issue.get("comment_count") is not None else "Falha",
            issue.get("person_comment_count") or 0,
            "Sim" if issue.get("person_commented") else "Não",
            safe_excel_text(url),
        ]
        sheet.append(row)
        current = sheet.max_row
        _link(sheet.cell(current, 1), url, result.get("base_url"))
        sheet.cell(current, 3).alignment = WRAP
        answer = sheet.cell(current, 13)
        if issue.get("person_commented"):
            answer.fill = YES_FILL
            answer.font = YES_FONT
        else:
            answer.fill = NO_FILL
            answer.font = NO_FONT
        for cell in sheet[current]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    _finish_sheet(sheet, [14, 48, 64, 22, 18, 22, 22, 32, 18, 20, 22, 34, 28, 54])


def _write_comments(sheet: Worksheet, result: dict) -> bool:
    _write_header(sheet, COMMENT_HEADERS)
    split = False
    for issue in result.get("issues") or []:
        url = str(issue.get("url") or "")
        for comment in issue.get("comments") or []:
            local = to_local(parse_azure_datetime(comment.get("created_at")))
            if local is None and comment.get("created_at_local"):
                local = parse_azure_datetime(comment.get("created_at_local"))
            chunks = text_chunks(safe_excel_text(html_to_plain_text(comment.get("text"))))
            if len(chunks) > 1:
                split = True
            for index, chunk in enumerate(chunks, start=1):
                sheet.append(
                    [
                        issue.get("id"),
                        safe_excel_text(issue.get("title")),
                        safe_excel_text(issue.get("assigned_to")),
                        safe_excel_text(comment.get("author")),
                        local.strftime("%d/%m/%Y") if local else "",
                        local.strftime("%H:%M") if local else "",
                        chunk,
                        safe_excel_text(url),
                        index if len(chunks) > 1 else 1,
                    ]
                )
                _link(sheet.cell(sheet.max_row, 1), url, result.get("base_url"))
                sheet.cell(sheet.max_row, 7).alignment = WRAP
    _finish_sheet(sheet, [14, 42, 28, 32, 18, 12, 80, 54, 10])
    return split


def _person_target_from_export(query: dict, result: dict) -> PersonTarget:
    person = result.get("person") or {}
    return PersonTarget(
        str(query.get("person_name") or person.get("display_name") or ""),
        person.get("identity_id") or query.get("person_identity_id"),
        person.get("unique_name") or query.get("person_unique_name"),
    )


def _write_full_discussions(sheet: Worksheet, result: dict, query: dict) -> bool:
    _write_header(sheet, FULL_DISCUSSION_HEADERS)
    person = _person_target_from_export(query, result)
    split = False
    issues = sorted(result.get("issues") or [], key=lambda item: int(item.get("id") or 0))
    for issue in issues:
        url = str(issue.get("url") or "")
        issue_has_person = "Sim" if issue.get("person_commented") else "Não"
        all_comments = list(issue.get("all_comments") or [])
        for comment in all_comments:
            local = to_local(parse_azure_datetime(comment.get("created_at")))
            if local is None and comment.get("created_at_local"):
                local = parse_azure_datetime(comment.get("created_at_local"))
            is_person = person_matches(comment, person)
            chunks = text_chunks(safe_excel_text(html_to_plain_text(comment.get("text"))))
            if len(chunks) > 1:
                split = True
            for chunk in chunks:
                sheet.append(
                    [
                        issue.get("id"),
                        safe_excel_text(issue.get("title")),
                        safe_excel_text(issue.get("assigned_to")),
                        local.strftime("%d/%m/%Y") if local else "",
                        local.strftime("%H:%M") if local else "",
                        safe_excel_text(comment.get("author")),
                        chunk,
                        issue_has_person,
                        "Sim" if is_person else "Não",
                        safe_excel_text(url),
                    ]
                )
                row = sheet.max_row
                _link(sheet.cell(row, 1), url, result.get("base_url"))
                sheet.cell(row, 7).alignment = WRAP
                _style_yes_no(sheet.cell(row, 8), issue_has_person == "Sim")
                _style_yes_no(sheet.cell(row, 9), is_person)
    _finish_sheet(sheet, [14, 42, 28, 18, 12, 32, 80, 36, 36, 54])
    return split


def _style_yes_no(cell, yes: bool) -> None:
    if yes:
        cell.fill = YES_FILL
        cell.font = YES_FONT
    else:
        cell.fill = NO_FILL
        cell.font = NO_FONT


def _write_validation(sheet: Worksheet, validation: dict, *, base_url: str) -> None:
    _write_header(sheet, VALIDATION_HEADERS)
    manual_by_id = {
        int(item["issue_id"]): item
        for item in (validation.get("manual_reviews") or [])
        if item.get("issue_id") is not None
    }
    for item in validation.get("sample", {}).get("items") or []:
        issue_id = item.get("issue_id")
        manual = manual_by_id.get(int(issue_id)) if issue_id is not None else None
        url = str(item.get("url") or "")
        sheet.append(
            [
                issue_id,
                safe_excel_text(url),
                item.get("original", {}).get("comment_count"),
                item.get("original", {}).get("person_comment_count"),
                (item.get("reextracted") or {}).get("comment_count"),
                (item.get("reextracted") or {}).get("person_comment_count"),
                item.get("auto_status"),
                manual.get("status") if manual else "",
                manual.get("real_comment_count") if manual else "",
                manual.get("real_person_comment_count") if manual else "",
                safe_excel_text(manual.get("divergence_type") if manual else item.get("divergence")),
                safe_excel_text(manual.get("notes") if manual else ""),
                manual.get("reviewed_at") if manual else validation.get("validated_at"),
                safe_excel_text(manual.get("reviewed_by") if manual else ""),
            ]
        )
        if url:
            _link(sheet.cell(sheet.max_row, 1), url, base_url)
    methodology = (validation.get("methodology") or {}).get("limitation") or ""
    sheet.append([])
    sheet.append(["Metodologia", safe_excel_text(methodology)])
    metrics = validation.get("metrics") or {}
    sheet.append(["Aviso", safe_excel_text(metrics.get("disclaimer"))])
    _finish_sheet(sheet, [12, 54, 18, 18, 18, 18, 22, 18, 18, 18, 28, 48, 22, 24])


def _write_parameters(
    sheet: Worksheet,
    query: dict,
    result: dict,
    *,
    split_text: bool,
    validation: dict | None = None,
    include_all_discussions: bool = False,
) -> None:
    generated = datetime.now(DISPLAY_TZ).strftime("%d/%m/%Y %H:%M")
    limitations = list(result.get("limitations") or [])
    if split_text:
        limitations.append(
            "Um ou mais comentários excederam 32.767 caracteres e foram divididos em linhas de continuação, sem descarte de texto."
        )
    if result.get("error_message"):
        limitations.append(str(result["error_message"]))
    rows = [
        ("Organização", query.get("organization")),
        ("Projeto", query.get("project")),
        ("Tipo de Work Item", query.get("work_item_type")),
        ("Critério de período", "Issues criadas no ano" if query.get("period_criterion") == "created" else "Issues alteradas no ano"),
        ("Ano", query.get("year")),
        ("Intervalo", f"01/01/{query.get('year')} até antes de 01/01/{int(query.get('year') or 0) + 1}"),
        ("Pessoa pesquisada", query.get("person_name")),
        ("Identificador da pessoa", query.get("person_identity_id") or result.get("person", {}).get("identity_id") or ""),
        ("Critério dos comentários", "Todo o histórico" if query.get("comment_scope") == "all" else "Apenas o período selecionado"),
        ("Data e hora da geração", f"{generated} (America/Cuiaba)"),
        ("Total de Issues processadas", result.get("issues_found")),
        ("Total de comentários analisados", result.get("comments_analyzed")),
        ("Total de comentários encontrados", result.get("comments_matched")),
        ("Situação da consulta", _situation(result)),
        ("Limitações ou erros parciais", "\n".join(limitations) if limitations else "Nenhuma"),
        (
            "Fonte das discussões",
            "Discussões extraídas dos registros de histórico disponíveis no Azure DevOps.",
        ),
    ]
    if include_all_discussions:
        rows.append(
            (
                "Exportação de discussões",
                "Aba Discussões completas — todos os autores disponíveis no histórico.",
            )
        )
        rows.append(("Linhas de discussão exportadas", count_stored_all_comments(result)))
    if validation:
        metrics = validation.get("metrics") or {}
        rows.extend(
            [
                ("Validação — cobertura WIQL consistente", "Sim" if metrics.get("coverage_consistent") else "Não"),
                ("Validação — amostra", metrics.get("sample_size")),
                ("Validação — divergência automática (amostra)", metrics.get("sample_auto_divergencia")),
                ("Validação — conferência manual (confere)", metrics.get("manual_confere")),
                ("Validação — conferência manual (divergência)", metrics.get("manual_divergencia")),
            ]
        )
    sheet.append(["Parâmetro", "Valor"])
    _style_header(sheet, 2)
    for label, value in rows:
        sheet.append([label, safe_excel_text(value)])
        sheet.cell(sheet.max_row, 2).alignment = WRAP
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{sheet.max_row}"
    sheet.column_dimensions["A"].width = 42
    sheet.column_dimensions["B"].width = 88


def _situation(result: dict) -> str:
    status = result.get("status")
    labels = {
        "complete": "Completa",
        "partial": "Incompleta — houve falhas em parte das Issues",
        "failed": "Falhou",
        "cancelled": "Cancelada — resultado incompleto",
        "ambiguous": "Incompleta — identidade ambígua, é preciso selecionar a pessoa",
        "running": "Em andamento",
    }
    return labels.get(str(status), str(status or ""))


def _write_header(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    _style_header(sheet, len(headers))


def _style_header(sheet: Worksheet, width: int) -> None:
    for column in range(1, width + 1):
        cell = sheet.cell(1, column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _finish_sheet(sheet: Worksheet, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    last_column = get_column_letter(len(widths))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 30


def _link(cell, url: str, base_url: str | None) -> None:
    if not url or not base_url:
        return
    if urlparse(url).netloc != urlparse(base_url).netloc:
        return
    cell.hyperlink = url
    cell.font = Font(color="185A4A", underline="single")


def _format_stamp(value: object, *, with_time: bool) -> str:
    parsed = to_local(parse_azure_datetime(str(value) if value else None))
    if parsed is None:
        return ""
    if with_time:
        return parsed.strftime("%d/%m/%Y %H:%M")
    return parsed.strftime("%d/%m/%Y")
